"""Read-only provenance reconciliation; never produces official benchmark scores.

Recorded historical labels are evidence about a past run, not replacement
decisions for the paper Judge. In particular, workbook common gold columns
must not be mistaken for the actual gold used by each model's historical run.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from dataclasses import asdict
import json
from math import fsum
from pathlib import Path
import re
import sys

if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from evaluation.judge import Judge
from evaluation.metrics import score
from evaluation.reproduce import DATASETS, MODELS, V4, load_baseline, write_csv
from evaluation.validation import ROOT, digest, file_hash, load_gold

COMPONENT_DATASET = {
    "ordinary_and_unanswerable_1200": "ordinary1190",
    "reasoning_refreshed_100": "reasoning_old100",
    "reasoning_incremental_100": "reasoning_hard100",
    "cross_pdf_first_400": "cross_old400",
    "cross_pdf_second_400": "cross_hard400",
}


def historical_text(value: str) -> str:
    """Reconstruct a historical *cache lookup*, never normalize a prediction."""
    return re.sub(r"\s+", " ", value.replace("\u00a0", " ")).strip()


def historical_cache_key(row: dict) -> str:
    parts = ("qwen36_calibrated_semantic_triclass_v4", row["judge_rules_sha256"],
             row["dataset_id"], row["question"], row["gold_answer"], row["pred_answer"])
    return digest("\n---\n".join(parts))


def read_workbook(path: Path) -> dict:
    import openpyxl
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    result = {}
    try:
        for sheet in book:
            rows = sheet.iter_rows(values_only=True)
            groups, fields = next(rows), next(rows)
            positions, group = {}, None
            for i, (header, field) in enumerate(zip(groups, fields)):
                if header in MODELS:
                    group = header
                positions[(group, field)] = i
            for row in rows:
                key = (sheet.title, str(row[fields.index("Unit ID")]), str(row[fields.index("QA ID")]))
                if key in result:
                    raise ValueError(f"Duplicate historical workbook identity: {key}")
                result[key] = {
                    model: {field: row[pos] for (owner, field), pos in positions.items() if owner == model}
                    for model in MODELS
                }
    finally:
        book.close()
    return result


def read_subjects(path: Path) -> dict:
    import openpyxl
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    result = {}
    try:
        for row in book["2200逐题明细"].iter_rows(min_row=2, values_only=True):
            key = (COMPONENT_DATASET[row[2]], str(row[4]), str(row[5]))
            if key in result:
                raise ValueError(f"Duplicate subject identity: {key}")
            result[key] = {"question": row[26], "answer": row[27], "discipline": row[6]}
    finally:
        book.close()
    return result


def exact_question_candidates(dataset: str, paper: str, question: str, records: dict) -> list:
    """Do not repair IDs unless exact document/question equality supports it."""
    return [list(key) for key, record in records.items()
            if key[:2] == (dataset, paper) and record["question"] == question]


def reconcile_membership(workbook: dict, subjects: dict, model: str, raw: dict) -> dict:
    """Describe different populations; never add outsiders to a fixed cohort."""
    recorded, selected, outsiders, absent = 0, 0, [], []
    for key, models in workbook.items():
        item = models[model]
        correct = item["Answer Correct"] == 1
        recorded += correct
        if key in subjects:
            selected += correct
        elif item["Present"] == 1:
            source = raw[key]
            outsiders.append({"dataset": key[0], "paper": key[1], "qa": key[2],
                "correct": correct, "question": source["question"],
                "reference": source["answer"],
                "exact_subject_candidates": exact_question_candidates(*key[:2], source["question"], subjects)})
    for key in subjects:
        if key not in workbook or workbook[key][model]["Present"] != 1:
            absent.append(list(key))
    return {"recorded_correct": recorded, "subject_join_correct": selected,
            "correct_outside_subject_cohort": recorded - selected,
            "contains_out_of_cohort_predictions": bool(outsiders),
            "outside_records": outsiders, "subject_items_without_prediction": absent}


def verify_raw_projection(raw: str, prediction) -> None:
    """Independent check of adapter extraction on legal rows; no permissive parser."""
    obj = json.loads(raw)
    if (set(obj) != {"answer_pre", "evidence_pages"}
            or obj["answer_pre"] != prediction.answer_pre
            or any(type(p) is not int for p in obj["evidence_pages"])
            or sorted(set(obj["evidence_pages"])) != prediction.evidence_pages):
        raise ValueError(f"Adapter changed model content: {prediction.qa_id}")


def snapshot_qa() -> dict:
    return {str(p.relative_to(ROOT)): file_hash(p) for p in sorted((ROOT / "data/qa").rglob("*.json"))}


def recorded_tables(workbook: dict, subjects: dict, raw_models: dict, paper: dict) -> list[dict]:
    """Reaggregate recorded item metrics, independently of the summary CSV."""
    comparisons = []
    for model, name in MODELS.items():
        records = [(key, models[model]) for key, models in workbook.items() if models[model]["Present"] == 1]
        correct = Counter()
        for key, item in records:
            task = ("Unanswerable" if raw_models[model][key]["answer"] == "Unanswerable" else "General") if key[0] == "ordinary1190" else ("Reasoning" if key[0].startswith("reasoning") else "Multi-Document")
            correct[task] += item["Answer Correct"] == 1
        table2 = {"All": 100 * sum(correct.values()) / 2200,
                  **{task: 100 * correct[task] / count for task, count in
                     (("General", 1000), ("Unanswerable", 200), ("Reasoning", 200), ("Multi-Document", 800))}}
        for field, metric, scale in (("E Precision", "E-Precision", 100), ("E Recall", "E-Recall", 100),
                                     ("E F1", "E-F1", 100), ("Pred Page Count", "A-Pages", 1)):
            table2[metric] = scale * fsum(float(item[field] or 0) for _, item in records) / 2200
        discipline_counts = Counter(r["discipline"] for r in subjects.values())
        discipline_correct = Counter()
        for key, row in subjects.items():
            discipline_correct[row["discipline"]] += key in workbook and workbook[key][model]["Answer Correct"] == 1
        table3 = {"All": 100 * sum(discipline_correct.values()) / len(subjects),
                  **{d: 100 * discipline_correct[d] / n for d, n in discipline_counts.items()}}
        for table, metrics in (("table2", table2), ("table3", table3)):
            for metric, expected in paper[table][name].items():
                actual = metrics[metric]
                comparisons.append({"table": table, "model": name, "metric": metric, "paper": expected,
                    "recorded_items_reaggregated": actual, "display_match": float(f"{actual:.2f}") == expected})
    return comparisons


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--rescore-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    args.rescore_dir = args.rescore_dir.resolve()
    out = args.output_dir.resolve()
    if out.is_relative_to(ROOT / "sxz") or out.is_relative_to(ROOT / "data/qa"):
        parser.error("Read-only input directory")
    out.mkdir(parents=True, exist_ok=False)
    qa_before = snapshot_qa()
    source_hashes = {}

    def track(path):
        source_hashes[str(path.relative_to(ROOT))] = file_hash(path)
        return path

    def read(path):
        return json.loads(track(path).read_text(encoding="utf-8"))

    history = ROOT / "sxz" / V4
    manifest = read(history / "evaluation_manifest.json")
    paper_reference = read(ROOT / "evaluation/paper_reference.json")
    if file_hash(ROOT / paper_reference["source"]) != paper_reference["paper_sha256"]:
        raise ValueError("Manuscript changed")
    cache = read(history / "qwen36_answer_judge_cache.json")
    for key, entry in cache.items():
        if key != historical_cache_key(entry) or entry["judge_rules_sha256"] != manifest["judge_rules_sha256"]:
            raise ValueError("Historical cache binding does not match its recorded inputs/prompt")
    lookup = {(v["dataset_id"], v["question"], v["gold_answer"], v["pred_answer"]): (k, v)
              for k, v in cache.items()}
    workbook = read_workbook(track(ROOT / "data/results/evaluations/detail_11models_5datasets.xlsx"))
    subjects = read_subjects(track(ROOT / "sxz/final_2200_classification_statistics.xlsx"))
    for relative in ("sxz/evaluate_11models_5datasets_calibrated_fixeddenom_seed42_gpu2345_parallel_v4.py",
                     "sxz/build_table3_by_subject.py", "evaluation/prompts/semantic_judge.txt",
                     "论文/ICLR2027_ScienceDoc/iclr2027_conference.tex",
                     "evaluation/audit_score_provenance.py"):
        track(ROOT / relative)
    copies = []
    raw_models = {}
    for model in MODELS:
        raw_models[model] = {}
        for dataset in DATASETS:
            path = ROOT / "data/results" / model / (dataset + ".json")
            data = read(path)
            # Resolve the historical root alias without changing any source path/file.
            original = Path(manifest["result_files"][model][dataset].replace("/data/czj/pku/", str(ROOT) + "/", 1))
            same = original.is_file() and file_hash(track(original)) == file_hash(path)
            copies.append({"model": model, "dataset": dataset, "same_bytes": same,
                           "copy": str(path.relative_to(ROOT)), "original": str(original.relative_to(ROOT))})
            for paper, document in data.items():
                for qa_id, qa in document["QA"].items():
                    raw_models[model][(dataset, str(paper), qa_id)] = qa
    print(f"Historical result copies: {sum(r['same_bytes'] for r in copies)}/{len(copies)}", flush=True)
    membership = {model: reconcile_membership(workbook, subjects, model, raw_models[model]) for model in MODELS}
    comparisons = recorded_tables(workbook, subjects, raw_models, paper_reference)
    gold, metadata = load_gold()
    legacy_index = {(g.legacy_paper_id, g.legacy_qa_id): g for g in gold.values()}
    if len(legacy_index) != len(gold):
        raise ValueError("Current legacy identity collision")
    model = "Qwen3-VL-8B"
    predictions, counts, _, _ = load_baseline(ROOT / "data/results", gold, model)
    config = read(args.rescore_dir / "judge_config.json")
    cache_path = track(args.rescore_dir / "judge_cache.jsonl")
    prior = read(args.rescore_dir / "report.json")
    new_judge = Judge(config, cache_path)
    replay = score(gold, predictions, new_judge, detailed=True)
    for field in ("total", "answer_metrics", "evidence_metrics", "discipline_breakdown", "audit_diagnostics"):
        if replay[field] != prior[field]:
            raise ValueError(f"Previous rescore no longer replays: {field}")
    if replay["audit_diagnostics"]["judge_failures_counted_zero"]:
        raise ValueError("Incomplete content-bound cache")
    rows, paired_jobs, groups = [], [], defaultdict(Counter)
    for row in replay["rows"]:
        qid = row["qa_id"]
        g = gold[qid]
        p = predictions.get(qid)
        if p is None:
            rows.append({"qa_id": qid, "task": g.task, "group": "missing", "old_correct": False,
                         "new_correct": False, "historical_cache_key": "", "historical_input_differences": ""})
            groups["missing"]["items"] += 1
            continue
        source = p.audit["historical_source"]
        key = (Path(source["file"]).stem, source["legacy_paper"], source["legacy_qa"])
        raw = raw_models[model][key]
        recorded = workbook[key][model]
        old_correct = recorded["Answer Correct"] == 1
        old = lookup.get((key[0], historical_text(raw["question"]), historical_text(raw["answer"]),
                          recorded["Pred Answer"] or ""))
        if recorded["Answer Judge Called"] == 1 and not old:
            raise ValueError(f"Historical Judge cache input not found: {qid}")
        if old and ((old[1]["answer_verdict"] == "correct") != old_correct):
            raise ValueError(f"Workbook/cache decision mismatch: {qid}")
        if p.status == "legal":
            verify_raw_projection(raw[source["field"]], p)
        differences = []
        group = p.status
        if p.status == "legal" and g.task == "Unanswerable":
            group = "deterministic_unanswerable"
        elif p.status == "legal":
            if old is None:
                group = "historical_judge_skipped"
            else:
                entry = old[1]
                differences = [name for name, a, b in (
                    ("question", entry["question"], g.question),
                    ("reference", entry["gold_answer"], g.answer),
                    ("prediction", entry["pred_answer"], p.answer_pre)) if a != b]
                group = "identical_judge_inputs" if not differences else "changed_judge_inputs"
                if differences:
                    paired_jobs.append({"qa_id": qid, "gold": asdict(g), "prediction": asdict(p),
                        "historical_cache_key": old[0], "historical_inputs": entry,
                        "current_decision": row["judge"]["decision"], "differences": differences})
        c = groups[group]
        c["items"] += 1
        c["old_correct"] += old_correct
        c["new_correct"] += row["answer_correct"]
        c["correct_to_incorrect"] += old_correct and not row["answer_correct"]
        c["incorrect_to_correct"] += not old_correct and row["answer_correct"]
        rows.append({"qa_id": qid, "task": g.task, "group": group, "old_correct": old_correct,
            "new_correct": row["answer_correct"], "historical_cache_key": old[0] if old else "",
            "historical_input_differences": ";".join(differences)})
    # Independent integer count aggregation, without score()'s aggregation helpers.
    manual_correct = sum(r["new_correct"] for r in rows)
    if manual_correct != replay["answer_metrics"]["All"]["correct"] or len(rows) != 2200:
        raise ValueError("Independent fixed-denominator aggregation mismatch")
    for key, metric in replay["answer_metrics"].items():
        selected = rows if key == "All" else [r for r in rows if r["task"] == key]
        if sum(r["new_correct"] for r in selected) * 100 / len(selected) != metric["accuracy"]:
            raise ValueError(f"Independent task aggregation mismatch: {key}")
    for discipline, metric in replay["discipline_breakdown"].items():
        selected = [r for r in rows if gold[r["qa_id"]].discipline == discipline]
        if sum(r["new_correct"] for r in selected) * 100 / len(selected) != metric["accuracy"]:
            raise ValueError(f"Independent discipline aggregation mismatch: {discipline}")
    claude = membership["Claude-sonncet-5"]
    current_records = {("ordinary1190", g.legacy_paper_id, g.legacy_qa_id): {"question": g.question}
                       for g in gold.values() if g.task in {"General", "Unanswerable"}}
    for item in claude["outside_records"]:
        if item["dataset"] == "ordinary1190":
            item["exact_current_candidates"] = exact_question_candidates(item["dataset"], item["paper"], item["question"], current_records)
    if qa_before != snapshot_qa():
        raise ValueError("Immutable QA changed")
    result = {"scope": "provenance reconciliation, not an alternative evaluation protocol",
        "official_reproduction_verified": False,
        "historical_copies_identical": sum(r["same_bytes"] for r in copies), "historical_copies_total": len(copies),
        "historical_cache_bindings_verified": len(cache), "qwen8b_groups": dict(groups),
        "historical_table_display_matches": {table: sum(r["display_match"] for r in comparisons if r["table"] == table)
                                             for table in ("table2", "table3")},
        "qwen8b_baseline_counts": dict(counts), "qwen8b_current_correct": manual_correct,
        "qwen8b_cache_replay_identical": True, "independent_task_discipline_aggregation_identical": True,
        "paired_jobs": len(paired_jobs), "membership": membership,
        "source_hashes": source_hashes, "qa_hashes": qa_before, "dataset_hashes": metadata["dataset_hashes"],
        "limitations": ["Past weight bytes/revision are not bound by the historical cache.",
            "Exact historical Judge inputs isolate input changes, not prompt versus checkpoint effects.",
            "A cached label is not human ground truth.", "Current-release diagnostic scores do not replace paper numbers."]}
    write_csv(out / "qwen8b_items.csv", rows)
    write_csv(out / "historical_copies.csv", copies)
    write_csv(out / "recorded_table_comparison.csv", comparisons)
    for name, value in (("audit.json", result), ("paired_jobs.json", paired_jobs)):
        (out / name).write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"qwen8b_groups": dict(groups), "paired_jobs": len(paired_jobs),
        "claude_outside_records": len(claude["outside_records"]),
        "claude_correct_outside": claude["correct_outside_subject_cohort"]}, indent=2), flush=True)
    return 0 if all(r["same_bytes"] for r in copies) else 2


if __name__ == "__main__":
    raise SystemExit(main())
