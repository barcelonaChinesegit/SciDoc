#!/usr/bin/env python3
"""Build the auditable classification workbook for the final 2,200 QA."""

from __future__ import annotations

import argparse
import json
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from pku_qa.workflows.selection.final_2200_contract import (
    CANONICAL_MODALITIES,
    CATEGORY_DEFINITIONS,
    CATEGORY_TAXONOMY,
    validate_final_dataset,
)


ROOT = Path(__file__).resolve().parents[4]
DEFAULT_MANIFEST = ROOT / "data/qa/7.final_2200/rel__collection__final_2200__manifest.json"
DEFAULT_OUTPUT = ROOT / "data/qa/7.final_2200/final_2200_classification_statistics.xlsx"

COMPONENT_LABELS = {
    "ordinary_and_unanswerable_1200": "普通与不可回答 1200",
    "reasoning_refreshed_100": "Reasoning 第一批 100",
    "reasoning_incremental_100": "Reasoning 第二批 100",
    "cross_pdf_first_400": "Cross-PDF 第一批 400",
    "cross_pdf_second_400": "Cross-PDF 第二批 400",
    "unanswerable": "不可回答 200",
    "reasoning_combined_200": "Reasoning 200",
    "cross_pdf_combined_800": "Cross-PDF 800",
    "ordinary_1000": "ordinary_qa.json",
    "unanswerable_200": "unanswerable_qa.json",
    "reasoning_200": "reasoning_qa.json",
    "cross_pdf_800": "cross_pdf_qa.json",
}
MODALITY_LABELS = {
    "text": "文本",
    "image": "图像",
    "table": "表格",
    "formula": "公式",
}

NAVY = "18324A"
BLUE = "DCEAF5"
PALE_BLUE = "EEF5FA"
ORANGE = "D97706"
PALE_ORANGE = "FEF3C7"
WHITE = "FFFFFF"
GRAY = "64748B"
GRID = "CBD5E1"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--check-only", action="store_true")
    return parser.parse_args()


def read_json(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError(f"Expected JSON object: {path}")
    return payload


def modality_source(component_id: str) -> str:
    if component_id in {"reasoning_refreshed_100", "reasoning_incremental_100"}:
        return "绑定源 QA 的 modal_types 有序并集"
    if component_id == "cross_pdf_second_400":
        return "Claude Sonnet 5 基于题目、答案和证据页图像的 API 复核"
    if component_id == "cross_pdf_combined_800":
        return "四文件交付中的 QA modal_types；第二批已完成证据页视觉 API 复核"
    return "原生 QA 标注，经别名和顺序规范化"


def source_document_count(component_id: str, qa: dict[str, Any]) -> int:
    if component_id in {"cross_pdf_first_400", "cross_pdf_combined_800", "cross_pdf_800"}:
        values = qa.get("evidence_source_docs", [])
        if isinstance(values, list) and values:
            return len(set(values))
    if component_id in {"cross_pdf_second_400", "cross_pdf_combined_800", "cross_pdf_800"}:
        value = qa.get("source_document_count", 0)
        return int(value) if isinstance(value, int) else 0
    return 1


def build_rows(manifest: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    index = 0
    for component in manifest["components"]:
        component_id = str(component["id"])
        path = Path(component["path"])
        dataset = read_json(path)
        validate_final_dataset(dataset, label=component_id)
        for paper_id, paper in dataset.items():
            for qa_id, qa in paper["QA"].items():
                index += 1
                modalities = list(qa["modal_types"])
                pages = list(qa["evidence_pages"])
                doc_count = source_document_count(component_id, qa)
                normalization = (
                    qa.get("annotation_provenance", {})
                    .get("final_2200_core_normalization", {})
                    if isinstance(qa.get("annotation_provenance"), dict)
                    else {}
                )
                modality_review = (
                    qa.get("annotation_provenance", {}).get("modality_api_review", {})
                    if isinstance(qa.get("annotation_provenance"), dict)
                    else {}
                )
                rows.append(
                    {
                        "序号": index,
                        "组件": COMPONENT_LABELS[component_id],
                        "组件ID": component_id,
                        "JSON文件": path.name,
                        "论文或合订本ID": str(paper_id),
                        "QA_ID": str(qa_id),
                        "一级领域": paper["primary_category"],
                        "二级领域": paper["secondary_category"],
                        "问题小类": qa["question_category"],
                        "原始问题小类": normalization.get("original_question_category", ""),
                        "问题类型": qa["question_type"],
                        "原始问题类型": normalization.get("original_question_type", ""),
                        "可回答性": "不可回答" if qa["answer"] == "Unanswerable" else "可回答",
                        "文档范围": "跨PDF" if component_id.startswith("cross_pdf") else "单PDF",
                        "来源文档数": doc_count,
                        "证据页数": len(pages),
                        "是否跨页": "跨页" if len(pages) > 1 else "未跨页",
                        "模态组合": "+".join(modalities),
                        "中文模态组合": "+".join(MODALITY_LABELS[value] for value in modalities),
                        "模态数": len(modalities),
                        "单多模态": "多模态" if len(modalities) > 1 else "单模态",
                        "模态标注依据": (
                            "Claude Sonnet 5 基于题目、答案和 PDF 页图像的逐题 API 复核"
                            if modality_review
                            else modality_source(component_id)
                        ),
                        "模态API复核版本": str(modality_review.get("version", "")),
                        "模态API置信度": modality_review.get("confidence", ""),
                        "复核前模态": "+".join(modality_review.get("previous_modal_types", [])),
                        "模态复核理由": str(modality_review.get("rationale", "")),
                        "evidence_pages": ",".join(str(value) for value in pages),
                        "reasoning_type": str(qa.get("reasoning_type", "")),
                        "reasoning_focus": str(qa.get("reasoning_focus", "")),
                        "source_qa_ids": ",".join(str(value) for value in qa.get("source_qa_ids", [])),
                        "问题": qa["question"],
                        "答案": qa["answer"],
                    }
                )
    if len(rows) != 2200:
        raise ValueError(f"Expected 2,200 rows, found {len(rows)}")
    return rows


def manifest_from_four_files(data_dir: Path) -> dict[str, Any]:
    specs = (
        ("ordinary_qa.json", "ordinary_1000", "ordinary_qa.json"),
        ("unanswerable_qa.json", "unanswerable_200", "unanswerable_qa.json"),
        ("reasoning_qa.json", "reasoning_200", "reasoning_qa.json"),
        ("cross_pdf_qa.json", "cross_pdf_800", "cross_pdf_qa.json"),
    )
    components = []
    for filename, component_id, label in specs:
        path = data_dir / filename
        dataset = read_json(path)
        validate_final_dataset(dataset, label=component_id)
        count = sum(len(paper["QA"]) for paper in dataset.values())
        components.append({"id": component_id, "path": str(path), "qa_count": count,
                           "paper_count": len(dataset), "label": label})
    if [c["qa_count"] for c in components] != [1000, 200, 200, 800]:
        raise ValueError("Four final JSON files do not have expected counts")
    return {"schema_version": 1, "collection_id": "final_2200", "components": components}


def style_sheet(ws: Any, *, freeze: str = "A2", filter_rows: bool = True) -> None:
    ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    if ws.max_row >= 1:
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(color=WHITE, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30
    thin = Side(style="thin", color=GRID)
    for row in ws.iter_rows():
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if filter_rows and ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def fit_columns(ws: Any, *, maximum: int = 42) -> None:
    for column in range(1, ws.max_column + 1):
        values = [str(ws.cell(row, column).value or "") for row in range(1, min(ws.max_row, 150) + 1)]
        width = min(maximum, max(10, max((len(value) for value in values), default=0) + 2))
        ws.column_dimensions[get_column_letter(column)].width = width


def add_counter_sheet(
    wb: Workbook,
    title: str,
    counter: Counter[Any],
    total: int,
    headers: tuple[str, str, str] = ("类型", "数量", "占比"),
) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for key, count in sorted(counter.items(), key=lambda item: (-item[1], str(item[0]))):
        ws.append([key, count, count / total])
    for cell in ws["C"][1:]:
        cell.number_format = "0.00%"
    style_sheet(ws)
    fit_columns(ws, maximum=55)


def add_component_modality_sheet(wb: Workbook, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("模态按组成")
    headers = ["组件", *[MODALITY_LABELS[value] for value in CANONICAL_MODALITIES], "单模态", "多模态", "合计"]
    ws.append(headers)
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[row["组件"]].append(row)
    components = list(dict.fromkeys(row["组件"] for row in rows))
    for component in components:
        items = grouped[component]
        ws.append(
            [
                component,
                *[sum(value in row["模态组合"].split("+") for row in items) for value in CANONICAL_MODALITIES],
                sum(row["单多模态"] == "单模态" for row in items),
                sum(row["单多模态"] == "多模态" for row in items),
                len(items),
            ]
        )
    ws.append(
        [
            "总计",
            *[sum(value in row["模态组合"].split("+") for row in rows) for value in CANONICAL_MODALITIES],
            sum(row["单多模态"] == "单模态" for row in rows),
            sum(row["单多模态"] == "多模态" for row in rows),
            len(rows),
        ]
    )
    for cell in ws[ws.max_row]:
        cell.fill = PatternFill("solid", fgColor=PALE_ORANGE)
        cell.font = Font(bold=True, color=ORANGE)
    style_sheet(ws)
    fit_columns(ws)


def add_overview_sheet(wb: Workbook, manifest: dict[str, Any], rows: list[dict[str, Any]]) -> None:
    ws = wb.active
    ws.title = "总览"
    ws.append(["最终 2200 条分类统计", "数值", "说明"])
    metrics = [
        ("QA 总数", len(rows), "四个最终交付 JSON"),
        ("可回答", sum(row["可回答性"] == "可回答" for row in rows), "answer != Unanswerable"),
        ("不可回答", sum(row["可回答性"] == "不可回答" for row in rows), "精确标签 Unanswerable"),
        ("单 PDF", sum(row["文档范围"] == "单PDF" for row in rows), "普通、不可回答和 Reasoning"),
        ("跨 PDF", sum(row["文档范围"] == "跨PDF" for row in rows), "两批 Cross-PDF"),
        ("单模态", sum(row["单多模态"] == "单模态" for row in rows), "modal_types 去重后为 1"),
        ("多模态", sum(row["单多模态"] == "多模态" for row in rows), "modal_types 去重后大于 1"),
        ("跨页", sum(row["是否跨页"] == "跨页" for row in rows), "evidence_pages 多于 1 页"),
        ("模态缺失", sum(not row["模态组合"] for row in rows), "统一 contract 要求为 0"),
        ("一级领域数", len({row["一级领域"] for row in rows}), "旧分类说明的 8 个一级领域"),
        ("二级领域数", len({row["二级领域"] for row in rows}), "实际覆盖的小领域"),
        ("问题小类数", len({row["问题小类"] for row in rows}), "旧分类说明的 21 个问题类别"),
    ]
    for item in metrics:
        ws.append(item)
    ws.append([])
    ws.append(["组件", "QA 数量", "论文/合订本数量"])
    for component in manifest["components"]:
        ws.append(
            [
                COMPONENT_LABELS[str(component["id"])],
                component["qa_count"],
                component["paper_count"],
            ]
        )
    style_sheet(ws, filter_rows=False)
    for cell in ws[14]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 65


def add_taxonomy_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("分类口径")
    ws.append(["层级", "一级领域/项目", "问题小类/取值", "定义或统计规则"])
    for primary, categories in CATEGORY_TAXONOMY.items():
        for category in dict.fromkeys(categories):
            ws.append(["问题分类", primary, category, CATEGORY_DEFINITIONS[category]])
    for modality in CANONICAL_MODALITIES:
        ws.append(["模态", "modal_types", modality, MODALITY_LABELS[modality]])
    ws.append(["模态规则", "Reasoning", "源 QA 并集", "按 source_qa_ids 回查 4211 原生 QA，合并全部所需模态。"])
    ws.append(["模态规则", "最终四文件", "逐题视觉 API 复核", "Claude Sonnet 5 同时读取题目、标准答案和 PDF 页图像，逐题判断回答实际需要的模态。"])
    ws.append(["跨页", "evidence_pages", "多于 1 页", "不可回答题证据页为空，统计为未跨页。"])
    ws.append(["单/多模态", "modal_types", "1 / >1", "按去重后的规范模态数量计算。"])
    style_sheet(ws)
    fit_columns(ws, maximum=80)


def add_modality_audit_sheet(wb: Workbook, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("四文件视觉复核")
    ws.append(["文件", "QA", "API复核", "文本", "图像", "表格", "公式", "单模态", "多模态"])
    for component in (
        "ordinary_qa.json",
        "unanswerable_qa.json",
        "reasoning_qa.json",
        "cross_pdf_qa.json",
    ):
        items = [row for row in rows if row["组件"] == component]
        occurrences = Counter(
            modality
            for row in items
            for modality in row["模态组合"].split("+")
            if modality
        )
        ws.append(
            [
                component,
                len(items),
                sum(bool(row["模态API复核版本"]) for row in items),
                *[occurrences[value] for value in CANONICAL_MODALITIES],
                sum(row["单多模态"] == "单模态" for row in items),
                sum(row["单多模态"] == "多模态" for row in items),
            ]
        )
    style_sheet(ws)
    fit_columns(ws, maximum=80)
    style_sheet(ws)
    fit_columns(ws, maximum=80)


def build_workbook(manifest: dict[str, Any], rows: list[dict[str, Any]]) -> Workbook:
    wb = Workbook()
    wb.properties.title = "最终 2200 条 QA 分类统计"
    wb.properties.subject = "学科、问题类别、模态、跨页和跨 PDF 分类"
    wb.properties.creator = "PKU Paper QA"
    wb.properties.created = datetime(2026, 9, 1, tzinfo=timezone.utc)
    add_overview_sheet(wb, manifest, rows)
    add_counter_sheet(wb, "组成分布", Counter(row["组件"] for row in rows), len(rows))
    add_counter_sheet(wb, "一级领域分布", Counter(row["一级领域"] for row in rows), len(rows))
    add_counter_sheet(wb, "二级领域分布", Counter(row["二级领域"] for row in rows), len(rows))
    add_counter_sheet(wb, "问题小类分布", Counter(row["问题小类"] for row in rows), len(rows))
    add_counter_sheet(wb, "问题类型分布", Counter(row["问题类型"] for row in rows), len(rows))
    add_counter_sheet(wb, "模态组合分布", Counter(row["模态组合"] for row in rows), len(rows))
    modality_occurrences = Counter(
        modality for row in rows for modality in row["模态组合"].split("+") if modality
    )
    add_counter_sheet(wb, "模态出现次数", modality_occurrences, len(rows), ("模态", "出现题数", "占全部 QA"))
    add_component_modality_sheet(wb, rows)
    add_counter_sheet(
        wb,
        "单多模态与跨页",
        Counter(f"{row['单多模态']} / {row['是否跨页']}" for row in rows),
        len(rows),
    )
    cross_rows = [row for row in rows if row["文档范围"] == "跨PDF"]
    add_counter_sheet(
        wb,
        "跨PDF文档数",
        Counter(f"{row['来源文档数']} 篇来源文档" for row in cross_rows),
        len(cross_rows),
    )
    reasoning_rows = [row for row in rows if row["reasoning_type"]]
    add_counter_sheet(
        wb,
        "推理关系分布",
        Counter(row["reasoning_type"] for row in reasoning_rows),
        len(reasoning_rows),
    )
    detail = wb.create_sheet("2200逐题明细")
    headers = list(rows[0])
    detail.append(headers)
    for row in rows:
        detail.append([row[header] for header in headers])
    style_sheet(detail)
    fit_columns(detail, maximum=34)
    for name in ("问题", "答案"):
        detail.column_dimensions[get_column_letter(headers.index(name) + 1)].width = 70
    detail.column_dimensions[get_column_letter(headers.index("模态标注依据") + 1)].width = 50
    add_taxonomy_sheet(wb)
    add_modality_audit_sheet(wb, rows)
    return wb


def check_workbook(path: Path) -> None:
    if not path.is_file():
        raise ValueError(f"Missing workbook: {path}")
    wb = load_workbook(path, read_only=True, data_only=True)
    required = {
        "总览",
        "模态组合分布",
        "模态出现次数",
        "2200逐题明细",
        "分类口径",
        "四文件视觉复核",
    }
    if not required.issubset(wb.sheetnames):
        raise ValueError(f"Workbook is missing sheets: {sorted(required - set(wb.sheetnames))}")
    if wb["2200逐题明细"].max_row != 2201:
        raise ValueError("Workbook detail sheet does not contain exactly 2,200 QA")


def main() -> int:
    args = parse_args()
    manifest = read_json(args.manifest)
    rows = build_rows(manifest)
    if args.check_only:
        check_workbook(args.output)
    else:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        workbook = build_workbook(manifest, rows)
        workbook.save(args.output)
        check_workbook(args.output)
    print(json.dumps({"status": "current", "qa_count": len(rows), "output": str(args.output)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
