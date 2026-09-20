#!/usr/bin/env python3
"""Build an auditable old/new pool for the later weak-GPT error screen."""

from __future__ import annotations

import argparse
import hashlib
import json
import time
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from pku_qa.workflows.selection.single_pdf_release_views import (
    COMBINED_RUNTIME_INPUT,
    write_runtime_inputs,
)


ROOT = Path(__file__).resolve().parents[4]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--challenge",
        type=Path,
        default=COMBINED_RUNTIME_INPUT,
    )
    parser.add_argument(
        "--cross",
        type=Path,
        default=ROOT / "data/qa/4.cross_pdf/challenge/rel__cross_pdf__challenge__batch01__n400__v1.json",
    )
    parser.add_argument(
        "--reasoning",
        type=Path,
        default=ROOT / "data/qa/3.reasoning/work__reasoning__historical_clean__batch00__n100.json",
    )
    parser.add_argument("--new-cross", type=Path, required=True)
    parser.add_argument("--new-reasoning", type=Path, required=True)
    parser.add_argument("--manual-gpt-workbook", type=Path, required=True)
    parser.add_argument("--evaluation-root", type=Path, default=ROOT / "data/results/evaluations")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--recommended-output", type=Path, required=True)
    parser.add_argument("--target", type=int, default=1000)
    return parser.parse_args()


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def atomic_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{time.time_ns()}.tmp")
    temporary.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    temporary.replace(path)


def grouped_items(path: Path, dataset: str, is_new: bool) -> list[dict[str, Any]]:
    payload = read_json(path)
    rows = []
    for paper_id, paper in payload.items():
        for qa_id, qa in paper.get("QA", {}).items():
            selection_id = f"{dataset}:{paper_id}/{qa_id}"
            rows.append(
                {
                    "selection_id": selection_id,
                    "dataset": dataset,
                    "paper_or_bundle_id": str(paper_id),
                    "qa_id": str(qa_id),
                    "source_dataset": str(path.resolve()),
                    "is_new": is_new,
                    "qa": qa,
                }
            )
    return rows


def manual_incorrect_ids(path: Path) -> set[tuple[str, str]]:
    workbook = load_workbook(path, read_only=False, data_only=True)
    # Challenge119 is a stride-10 manual sample from the 1190-item Challenge
    # input.  Its incorrect item IDs are deliberately mapped back onto the
    # full ``challenge`` dataset when only the hardness score is built.
    mapping = {
        "Challenge119": "challenge",
        "Reasoning100": "reasoning",
        "CrossPDF100": "cross_pdf",
    }
    result: set[tuple[str, str]] = set()
    for sheet_name, dataset in mapping.items():
        sheet = workbook[sheet_name]
        header = [cell.value for cell in sheet[1]]
        item_index = header.index("item_id")
        correct_index = header.index("answer_correct")
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[item_index] and row[correct_index] is False:
                result.add((dataset, str(row[item_index])))
    return result


def judge_wrong_ids(path: Path, dataset: str) -> set[tuple[str, str]]:
    if not path.is_file():
        return set()
    payload = read_json(path)
    result = set()
    for paper_id, paper in payload.items():
        for qa_id, qa in paper.get("QA", {}).items():
            if qa.get("answer_is_correct") is False:
                result.add((dataset, f"{paper_id}/{qa_id}"))
    return result


def score_rows(
    rows: list[dict[str, Any]],
    manual_wrong: set[tuple[str, str]],
    local_wrong_by_model: dict[str, set[tuple[str, str]]],
) -> None:
    for row in rows:
        item_id = f"{row['paper_or_bundle_id']}/{row['qa_id']}"
        key = (row["dataset"], item_id)
        score = 100 if row["is_new"] else 0
        reasons = ["new_dual_reviewed_expansion"] if row["is_new"] else []
        if key in manual_wrong:
            score += 100
            reasons.append("gpt5_6_manual_incorrect")
        for model, wrong in local_wrong_by_model.items():
            if key in wrong:
                score += 25
                reasons.append(f"local_qwen_{model}_incorrect")
        qa = row["qa"]
        if qa.get("question_type") == "Inferential":
            score += 5
            reasons.append("inferential")
        if int(qa.get("source_document_count", 0) or 0) == 3:
            score += 8
            reasons.append("three_document")
        evidence_pages = qa.get("evidence_pages")
        if isinstance(evidence_pages, list) and len(evidence_pages) >= 2:
            try:
                if max(evidence_pages) - min(evidence_pages) >= 3:
                    score += 3
                    reasons.append("nonlocal_evidence")
            except TypeError:
                pass
        row["hardness_priority_score"] = score
        row["hardness_priority_reasons"] = reasons


def select_recommended(rows: list[dict[str, Any]], target: int) -> list[dict[str, Any]]:
    new_rows = [row for row in rows if row["is_new"]]
    if len(new_rows) > target:
        raise ValueError(f"new rows exceed target: {len(new_rows)} > {target}")
    chosen = sorted(
        new_rows,
        key=lambda row: (-row["hardness_priority_score"], row["selection_id"]),
    )
    selected_ids = {row["selection_id"] for row in chosen}
    old_rows = sorted(
        (row for row in rows if row["selection_id"] not in selected_ids),
        key=lambda row: (-row["hardness_priority_score"], row["selection_id"]),
    )
    chosen.extend(old_rows[: max(0, target - len(chosen))])
    if len(chosen) != target:
        raise ValueError(f"insufficient selection pool: {len(chosen)}/{target}")
    return chosen


def payload(rows: list[dict[str, Any]], status: str) -> dict[str, Any]:
    fingerprint = hashlib.sha256(
        "\n".join(row["selection_id"] for row in rows).encode("utf-8")
    ).hexdigest()
    return {
        "status": status,
        "count": len(rows),
        "dataset_distribution": dict(sorted(Counter(row["dataset"] for row in rows).items())),
        "new_item_count": sum(row["is_new"] for row in rows),
        "known_gpt_incorrect_count": sum(
            "gpt5_6_manual_incorrect" in row["hardness_priority_reasons"] for row in rows
        ),
        "selection_id_sha256": fingerprint,
        "caveat": (
            "This is a hardness-prioritized preselection, not evidence that GPT "
            "accuracy is below 80%. The collaborator's weak-GPT screen must make "
            "the final error-based selection."
        ),
        "generated_at_utc": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
        "items": rows,
    }


def main() -> None:
    args = parse_args()
    if args.challenge.resolve() == COMBINED_RUNTIME_INPUT.resolve():
        write_runtime_inputs()
    rows = []
    rows.extend(grouped_items(args.challenge, "challenge", False))
    rows.extend(grouped_items(args.cross, "cross_pdf", False))
    rows.extend(grouped_items(args.reasoning, "reasoning", False))
    rows.extend(grouped_items(args.new_cross, "cross_pdf_new", True))
    rows.extend(grouped_items(args.new_reasoning, "reasoning_new", True))
    manual_wrong = manual_incorrect_ids(args.manual_gpt_workbook)
    eval_paths = {
        "4b": {
            "challenge": args.evaluation_root / "single_pdf_1200/full_1200/judge_4B.json",
            "cross_pdf": args.evaluation_root / "cross_pdf_400_full_pdf/judge_4B.json",
            "reasoning": args.evaluation_root / "reasoning_qa_100_full_pdf/judge_4B.json",
        },
        "8b": {
            "challenge": args.evaluation_root / "single_pdf_1200/full_1200/judge_8B.json",
            "cross_pdf": args.evaluation_root / "cross_pdf_400_full_pdf/judge_8B.json",
            "reasoning": args.evaluation_root / "reasoning_qa_100_full_pdf/judge_8B.json",
        },
    }
    wrong_by_model = {
        model: set().union(
            *(judge_wrong_ids(path, dataset) for dataset, path in paths.items())
        )
        for model, paths in eval_paths.items()
    }
    score_rows(rows, manual_wrong, wrong_by_model)
    rows.sort(key=lambda row: (-row["hardness_priority_score"], row["selection_id"]))
    recommended = select_recommended(rows, args.target)
    atomic_json(args.output, payload(rows, "screening_pool"))
    atomic_json(args.recommended_output, payload(recommended, "preselection_pending_weak_gpt"))


if __name__ == "__main__":
    main()
